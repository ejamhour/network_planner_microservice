from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt


REFERENCE = Path(
    r"D:\CISEI\TELECOM\PlanApp\Relatorios\Relatorio_06_2026_revisado.docx"
)
OUTPUT = Path(
    r"D:\CISEI\TELECOM\microservices\network-service\network_planner_microservice\Relatorio_08_2026_atividades.docx"
)


def clear_body_keep_section(doc: Document) -> None:
    body = doc._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)
    if sect_pr is not None and body.sectPr is None:
        body.append(sect_pr)


def add_paragraph(doc: Document, text: str, style: str = "Normal"):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    return paragraph


def add_blank(doc: Document):
    return doc.add_paragraph("")


def main() -> None:
    doc = Document(REFERENCE)
    clear_body_keep_section(doc)

    add_paragraph(doc, "Relatório de Atividades (AGO/26)", "Heading 1")
    add_paragraph(doc, "Edgard Jamhour", "Normal")
    add_blank(doc)
    add_paragraph(
        doc,
        "As atividades deste período estiveram concentradas na implementação do microserviço de planejamento de redes. Esse microserviço utiliza o microserviço de features geográficas, já consolidado, como base para a extração de características ambientais dos enlaces. A nova camada em desenvolvimento transforma essas características, juntamente com dados de equipamentos, antenas e interfaces, em grafos de conectividade e resultados de planejamento.",
        "Normal",
    )
    add_paragraph(
        doc,
        "Dessa forma, o trabalho descrito neste relatório representa uma etapa diretamente ligada ao produto final. O objetivo não foi apenas realizar experimentos isolados, mas construir a arquitetura que permitirá oferecer planejamento de redes por API, notebooks e interfaces assistidas por linguagem natural, reaproveitando o serviço de features como componente especializado dentro de uma solução mais ampla.",
        "Normal",
    )
    add_blank(doc)

    sections = [
        (
            "1. Reestruturação das classes de planejamento",
            [
                "Neste período, as atividades concentraram-se na reorganização da base de planejamento de redes, com ênfase na separação entre o algoritmo RPL, a preparação do grafo de candidatos e a representação física dos elementos de rede. Essa separação tornou o sistema mais adequado para tratar cenários diferentes, como redes celulares, redes Wi-SUN, enlaces com repetidores e soluções híbridas com múltiplas tecnologias.",
                "Foi revisada a nomenclatura das classes de planejamento para refletir melhor a modelagem adotada. A estrutura passou a distinguir site, dispositivo e interface de rádio. O site representa a posição física; o dispositivo representa o equipamento instalado; e a interface representa o ponto efetivo de conexão usado pelo RPL. Essa decisão resolveu ambiguidades importantes, principalmente nos casos em que um mesmo local possui mais de uma interface ou tecnologia.",
                "Também foi consolidado o tratamento de coordenadas geográficas e UTM. O usuário pode fornecer posições em latitude e longitude ou em coordenadas projetadas, desde que os dados estejam consistentes. O sistema passou a admitir uma projeção de trabalho explícita, reduzindo recalculações e evitando misturas indevidas de sistemas de referência.",
            ],
        ),
        (
            "2. Criação do GraphPlanner e configuração por perfis",
            [
                "Foi desenvolvido um planejador genérico responsável por transformar registros preparados pelo usuário em objetos de planejamento, interfaces RPL e grafos de candidatos. A leitura de arquivos permanece fora da classe, permitindo que notebooks, APIs ou outros módulos forneçam os dados em estruturas Python padronizadas. Essa decisão evita dependência direta de CSV e preserva a possibilidade de uso futuro via JSON ou chamadas de API.",
                "O GraphPlanner não é apenas uma etapa auxiliar do planejamento celular. Ele foi concebido como uma camada genérica de modelagem de rede, capaz de representar sites com dispositivos, dispositivos com múltiplas interfaces e interfaces com diferentes tecnologias, frequências, potências e antenas. Para o RPL, cada interface pode ser tratada como um nó do grafo, enquanto o GraphPlanner preserva a relação entre interface, dispositivo e site físico.",
                "Essa estrutura permite representar cenários em que um mesmo equipamento possui, por exemplo, uma interface LTE para comunicação com torre e outra interface para comunicação nó-a-nó. Também permite criar enlaces internos entre interfaces de um mesmo dispositivo, quando o equipamento é capaz de rotear entre tecnologias. Essa modelagem é essencial para redes híbridas, nas quais a solução final pode combinar acesso celular, enlaces locais e múltiplos saltos controlados.",
                "A configuração passou a ser definida por perfis em TOML. Os perfis descrevem antenas, interfaces, dispositivos, regras de conectividade e mapeamento de métricas por tecnologia. Essa abordagem reduz a quantidade de código necessária nos notebooks e cria uma base mais adequada para uso futuro por API ou por uma interface assistida por LLM.",
                "O GraphPlanner também passou a produzir relatórios de configuração e tabelas de arestas, permitindo ao usuário inspecionar como os sites, dispositivos, interfaces e candidatos foram definidos antes de executar cálculos mais custosos. Essa etapa é importante porque erros de configuração devem ser identificados antes da extração de características geográficas.",
            ],
        ),
        (
            "3. Métricas de enlace e integração com o serviço geográfico",
            [
                "O compilador de métricas foi revisado para trabalhar com informações estruturadas por enlace. Além das características ambientais extraídas do serviço geográfico, o contexto da métrica passou a incluir dados do transmissor e do receptor, como potência de transmissão, ganho de antena, altura de instalação, frequência e tipo de antena.",
                "Foi criada uma métrica específica para enlaces celulares entre torre e nó, combinando obstruções de terreno e edifícios de forma simplificada. Para edifícios, a métrica considera invasões nas zonas de Fresnel e no núcleo do enlace, com compressão logarítmica da profundidade para evitar penalizações excessivas em situações onde múltiplos obstáculos podem inflar artificialmente a medida acumulada.",
                "A integração assíncrona com o serviço geográfico também foi reforçada. Foram adicionadas rotinas de repetição para respostas 503, associadas a indisponibilidade temporária de workers, e a chave de cache passou a considerar parâmetros relevantes do enlace, como alturas e frequência. Isso melhora a estabilidade dos testes em lote e evita recomputações desnecessárias.",
            ],
        ),
        (
            "4. Planejamento celular com setores",
            [
                "Foi criada a primeira especialização sobre o GraphPlanner, denominada CellPlanner. Essa classe trata o caso de planejamento baseado em sites conectados, como torres ou pontos já integrados ao backbone, oferecendo candidatos de conexão direta para nós clientes usando uma tecnologia primária.",
                "O CellPlanner introduziu a seleção geométrica de candidatos considerando antenas setoriais. Antes de chamar o serviço geográfico, o planejador verifica se o cliente está dentro do azimute e da abertura angular do setor. Com isso, enlaces claramente inviáveis deixam de ser avaliados, reduzindo custo computacional e facilitando a análise visual do grafo candidato.",
                "Também foram adicionados métodos auxiliares para sugerir rotação de setores e identificar setores vazios. A rotação usa uma função simples baseada na fração de clientes descobertos e no desbalanceamento de carga entre setores. Setores sem candidatos não são penalizados automaticamente na função de rotação; eles são apenas reportados e podem ser removidos por um método específico.",
                "A estrutura atual permite que o usuário ajuste manualmente perfis de torre, por exemplo usando três setores em uma torre e dois setores em outra. Esse recurso será importante para estudos de cobertura fixa, nos quais poucos pontos podem estar mal distribuídos em relação aos setores originalmente definidos.",
            ],
        ),
        (
            "5. Validação em notebooks e operação dos serviços",
            [
                "Foram criados e revisados notebooks de demonstração para validar o uso das novas classes. Os exemplos passaram de métricas artificiais para métricas calculadas pelo serviço geográfico, usando pontos reais carregados a partir de CSV. Também foram testados cenários com uma interface e com múltiplas interfaces, permitindo observar a formação do grafo candidato e o resultado do RPL.",
                "Nos testes de múltiplas interfaces, alguns nós foram configurados com uma interface LTE e uma interface nó-a-nó, usando métricas diferentes por tecnologia. O grafo passou a conter enlaces de acesso, enlaces locais e enlaces internos entre interfaces do mesmo dispositivo. Esse teste validou a hipótese central da arquitetura: o RPL pode permanecer simples e agnóstico, desde que receba um grafo previamente preparado com as interfaces e custos corretos.",
                "Essa validação é relevante porque mostra que o GraphPlanner já suporta problemas mais amplos do que a conexão direta entre torre e cliente. A mesma base pode ser usada para planejamento LTE puro, LTE com recuperação por enlaces nó-a-nó, Wi-SUN com nós repetidores e, posteriormente, redes híbridas com restrições específicas definidas por classes especializadas.",
                "O notebook do CellPlanner foi simplificado para deixar explícita a sequência de execução: carregar pontos, criar o planejador, construir candidatos, inspecionar relatórios opcionais, calcular métricas e executar o RPL. Relatórios e visualizações pré-métrica foram agrupados em uma única célula opcional, evitando que mutações como remoção de setores limpem o grafo de candidatos de forma inesperada.",
                "Além do desenvolvimento de código, foram avaliadas alternativas para executar o serviço de planejamento fora do VS Code. Foi identificado o comando de inicialização do FastAPI, o uso do hub em app.main_hub:hub_app e as variáveis de ambiente necessárias para acesso ao MinIO e ao cache local. Também foram discutidos ajustes de memória do WSL para reduzir pressão sobre o sistema durante execuções com Docker e processamento geográfico.",
            ],
        ),
        (
            "6. Conclusão do período",
            [
                "O período consolidou uma mudança importante na arquitetura do planejador. O RPL permaneceu agnóstico, operando apenas sobre interfaces e custos de arestas, enquanto a responsabilidade de construir grafos candidatos passou para classes de planejamento. Essa separação melhora a reutilização do algoritmo e permite criar especializações sem modificar o núcleo do RPL.",
                "A combinação de perfis TOML, GraphPlanner, métricas por tecnologia e suporte a múltiplas interfaces cria uma base flexível para os próximos cenários. O CellPlanner passa a ser uma especialização construída sobre essa base, e não o centro da arquitetura. Os próximos passos naturais são implementar o planejamento celular com recuperação por enlaces nó-a-nó, criar uma especialização para redes Wi-SUN e adicionar exportações mais completas, como GeoJSON, para integração com ferramentas externas.",
                "Os resultados obtidos indicam que a abordagem é promissora para planejamento de infraestrutura rural e urbana, especialmente em cenários de automação, smart grid e redes híbridas. A arquitetura atual ainda exige calibração de métricas e refinamento das especializações, mas já oferece uma base operacional para o microserviço de planejamento, que passa a consumir o microserviço de features e a expor resultados em formatos adequados para notebooks, APIs e futuras interfaces assistidas por linguagem natural.",
            ],
        ),
    ]

    for heading, paragraphs in sections:
        add_paragraph(doc, heading, "Heading 2")
        for paragraph in paragraphs:
            add_paragraph(doc, paragraph, "Normal")
        add_blank(doc)

    # Preserve source styles but make the output metadata current.
    props = doc.core_properties
    props.title = "Relatório de Atividades (AGO/26)"
    props.author = "Edgard Jamhour"
    props.subject = "Atividades de desenvolvimento do planejador de redes"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
